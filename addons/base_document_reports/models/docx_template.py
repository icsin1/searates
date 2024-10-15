# -*- coding:utf-8 -*-
import io
import re
import mimetypes
import base64
import tempfile

from docxtpl import InlineImage
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from odoo.tools.misc import formatLang
from odoo import api, fields, models, _
from odoo.tools.mimetypes import get_extension
from odoo.tools.misc import xlsxwriter
from odoo.exceptions import ValidationError
from odoo.tools.misc import format_datetime, format_date, get_lang

IGNORE_MODELS = ['res.users']
IGNORE_MODEL_FIELDS = ['__last_update', 'write_uid', 'create_uid', 'product_tmpl_id', 'product_variant_id']
IGNORE_FIELD_KEYS = ['my_activity_', 'activity_', 'message_', 'website_message_']
COMPANY_ALLOWED_FIELDS = [
    'name', 'email', 'company_registry', 'website', 'favicon', 'logo', 'logo_web', 'mobile', 'phone', 'state_id', 'zip', 'vat', 'city', 'street', 'street2', 'logo_1024', 'logo_512', 'logo_256',
    'logo_128', ]


class DocxTemplate(models.Model):
    _name = "docx.template"
    _inherit = 'mixin.report.template'
    _description = "Document Report Template"
    _template_field = 'docx_file'
    _report_type = 'docx'

    docx_file = fields.Binary(string='Template (.docx)', required=True)
    filename = fields.Char(string='Filename')
    company_ids = fields.Many2many('res.company', 'res_company_docx_template', string='Allowed Companies')
    view_type = fields.Selection(selection_add=[
        ('list', 'List'),
        ('form', 'Form'),
        ('both', 'Both'),
    ], default='form', string='Data Representation', required=True, ondelete={'both': 'cascade'})
    output_type = fields.Selection(selection_add=[
        ('pdf', 'PDF'),
        ('docx', 'Original')
    ], default='pdf', string='Output Type', required=True, ondelete={'docx': 'cascade', 'pdf': 'cascade'})
    lazy_load = fields.Boolean(default=False, help='If Using Lazy Load, use obj.technical_field (record set chain)')
    model_name = fields.Char(string='Model Name', related='res_model_id.model', store=True)
    usable_variable_file = fields.Binary(string='Document Fields')
    usable_variable_filename = fields.Char()
    usable_variable_technical_file = fields.Binary('Document Technical Fields')
    usable_variable_technical_filename = fields.Char()
    terms_and_conditions = fields.Html(string='Default Terms & Conditions')
    active = fields.Boolean(default="True")

    @api.constrains('docx_file')
    def _check_file_name(self):
        for template in self:
            mimetype = mimetypes.guess_type(template.filename)[0]
            has_extension = get_extension(template.filename) or mimetype
            if has_extension != '.docx':
                raise ValidationError(_('Allowed Format Docx'))

    def _html_to_text(self, html_value):
        # Removing <br> with \n
        html_value = str(html_value or '').replace("<br>", "\n")
        # and removing all HTML tags
        return re.sub(r'<[^<]+?>', '', html_value)

    def convert_to_dict(self, record, ignore_fields=[], allowed_fields=[], only_record_fields=False, docx_template=False):
        fields = record._fields
        if allowed_fields:
            fields = {field: fields[field] for field in fields if field in allowed_fields}
        ignore_fields += IGNORE_MODEL_FIELDS
        values = {}
        for field in fields:
            if not any([field.startswith(ignore_key) for ignore_key in IGNORE_FIELD_KEYS]) and field not in ignore_fields:
                field_type = fields[field].type
                record_value = record[field]

                # For Image File Convert value to InlineImage Object to be able to show image in file
                if field_type == 'binary' and docx_template:
                    if record_value:
                        # FIXME: Check with File streaming
                        tmp_img_file = tempfile.NamedTemporaryFile(delete=None, suffix='.jpg')
                        tmp_img_file.write(base64.urlsafe_b64decode(record_value))
                        record_value = InlineImage(docx_template, tmp_img_file.name)
                        tmp_img_file.close()
                elif field_type == 'html':
                    if record_value:
                        record_value = self._html_to_text(record[field])
                elif field_type == 'monetary':
                    record_value = formatLang(self.env, record_value, currency_obj=record[fields[field].get_currency_field(record)] or self.env.company.currency_id)
                elif field_type == 'selection':
                    # Show selection field Label-value instead of Key
                    record_value = dict(record._fields[field]._description_selection(self.env)).get(record_value)
                else:
                    if record_value and field_type == 'date':
                        record_value = format_date(self.env, record_value, lang_code=get_lang(self.env).code, date_format=False)
                    if record_value and field_type == 'datetime':
                        tz = self.env.user.tz
                        if (self.env.user._is_public()) and getattr(record, 'company_id', None) is not None:
                            tz = record.company_id.tz
                        record_value = format_datetime(self.env, record_value, tz=tz, lang_code=get_lang(self.env).code, dt_format=False)

                if field_type in ['one2many']:
                    if not only_record_fields:
                        o2m_ignore_fields = [fields[field].inverse_name, field]
                        values[self._key_to_camel_case(field, ignore_name=True)] = record_value and [
                            self.convert_to_dict(rec, ignore_fields=o2m_ignore_fields, docx_template=docx_template) for rec in record_value
                        ]
                        values.update({**self._convert_to_dict_dynamic_rows(record, field, field_type, record_value)})
                    else:
                        values[self._key_to_camel_case(field)] = f'{len(record_value)} records'

                elif field_type == 'many2one':
                    m2o_key = self._key_to_camel_case(field, ignore_name=not only_record_fields)
                    # For Company many2one field show only allowed fields
                    if fields[field].comodel_name in IGNORE_MODELS or only_record_fields:
                        values[m2o_key] = record_value.display_name if record_value and record_value.check_access_rights('read', raise_exception=False) else ''
                    elif record_value and fields[field].comodel_name == 'res.company':
                        values[m2o_key] = self.convert_to_dict(record_value, allowed_fields=COMPANY_ALLOWED_FIELDS, only_record_fields=True, docx_template=docx_template)
                    else:
                        values[m2o_key] = self.convert_to_dict(record_value, only_record_fields=True, docx_template=docx_template)
                elif not only_record_fields and field_type == 'many2many':
                    values[self._key_to_camel_case(field, ignore_name=True)] = record_value
                else:
                    values[self._key_to_camel_case(field)] = record_value if record_value else ''
        return values

    def _convert_to_dict_dynamic_rows(self, record, field, field_type, record_value):
        return {}

    def _key_to_camel_case(self, field_key, ignore_name=False):
        return ''.join([key.capitalize() for key in field_key.replace("_id", "_name" if not ignore_name else '').split('_')])

    def _add_dynamic_field_row(self, sheet, field_name, description, row_no, add_technical_name=False):
        sheet.write(row_no, 0, field_name)
        sheet.write(row_no, 1, description)
        if add_technical_name:
            sheet.write(row_no, 2, field_name)
        # returning added field name and description
        return field_name, description

    def _add_field_row(self, sheet, field, row_no, add_technical_name=False, parent_field=''):
        self.ensure_one()
        is_m2m_o2m = field.ttype in ['one2many', 'many2many']
        template_field_name = self._key_to_camel_case(field.name, ignore_name=is_m2m_o2m)
        field_name = field.name
        description = field.field_description
        if parent_field:
            template_field_name = '%s.%s' % (parent_field[0], template_field_name)
            field_name = '%s.%s' % (parent_field[0], field_name)
            description = '%s > %s' % (parent_field[1], description)
        if is_m2m_o2m:
            description = '%s (Array) - See %s Sheet' % (description, template_field_name)
        sheet.write(row_no, 0, template_field_name)
        sheet.write(row_no, 1, description)
        if add_technical_name:
            sheet.write(row_no, 2, field_name)
        # returning added field name and description
        return template_field_name, description

    def _filter_fields(self, fields, ignore_x2o=False):
        fields = fields.filtered(lambda f: not any([f.name.startswith(ignore_key) for ignore_key in IGNORE_FIELD_KEYS]))
        fields = fields.filtered(lambda f: not any([f.name.startswith(ignore_key) for ignore_key in IGNORE_MODEL_FIELDS]))
        if ignore_x2o:
            fields = fields.filtered(lambda f: f.ttype not in ['many2many', 'one2many'])
        return fields

    def download_usable_variable(self):
        self.ensure_one()
        return self._generate_usable_variable(download_file=True)

    def download_usable_variable_technical(self):
        return self._generate_usable_variable(add_technical_name=True, download_file=True)

    def _generate_usable_variable(self, model=False, worksheet=False, parent_field_obj=None, model_type=None, relation_field=None, add_technical_name=False, download_file=False):
        self.ensure_one()
        self = self.sudo()
        model = model or self.res_model_id

        fields = model.field_id.sorted(lambda f: f.ttype in ['many2one', 'one2many', 'many2many'])
        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})

        heading_style = workbook.add_format({'align': 'center', 'bold': True, 'font_size': '14px'})
        if not worksheet:
            worksheet = workbook.add_worksheet(model.name)
        merge_column_number = 1
        if add_technical_name:
            merge_column_number = 2
        worksheet.merge_range(0, 0, 0, merge_column_number, model.name, heading_style)

        worksheet.write(1, 0, 'Template Field', heading_style)
        worksheet.set_column(1, 0, 30)

        worksheet.write(1, 1, 'Description', heading_style)
        worksheet.set_column(1, 1, 30)

        if add_technical_name:
            worksheet.write(1, 2, 'Technical Field', heading_style)
            worksheet.set_column(1, 2, 30)

        row = 2
        for field in self._filter_fields(fields, ignore_x2o=bool(parent_field_obj)):
            field_type = field.ttype
            row += 1

            if field_type == 'many2one' and field.relation not in IGNORE_MODELS and field.name != relation_field:
                # adding many2one fields
                parent_field = self._key_to_camel_case(field.name, ignore_name=True), field.field_description
                m2o_model = self.env['ir.model'].sudo().search([('model', '=', field.relation)], limit=1)
                fields = m2o_model.field_id

                # For Company model use only allowed company fields
                if field.relation == 'res.company':
                    fields = fields.filtered(lambda f: f.name in COMPANY_ALLOWED_FIELDS)

                for m2o_field in self._filter_fields(fields, ignore_x2o=True):
                    row += 1
                    self._add_field_row(worksheet, m2o_field, row, parent_field=parent_field, add_technical_name=add_technical_name)

            elif field_type in ['one2many', 'many2many'] and field.relation not in IGNORE_MODELS:
                # adding one2many and many2many first level
                parent_field = self._add_field_row(worksheet, field, row, add_technical_name)
                x2o_model = self.env['ir.model'].sudo().search([('model', '=', field.relation)], limit=1)
                sheet = workbook.add_worksheet(parent_field[0])
                self._generate_usable_variable(x2o_model, sheet, model_type=field_type, parent_field_obj=field, relation_field=field.relation_field, add_technical_name=add_technical_name)
                # Used to generated dynamic fields for x2m
                row = self._generate_dynamic_fields(row, x2o_model, worksheet, field_type, field, add_technical_name=add_technical_name)
            else:
                self._add_field_row(worksheet, field, row, add_technical_name)

        workbook.close()

        file_name = '%s%s-TemplateVariables.xlsx' % (('Technical-' if add_technical_name else ''), self.name)
        xlsx_base64 = base64.b64encode(output.getvalue())

        self.write({
            'usable_variable_technical_filename' if add_technical_name else 'usable_variable_filename': file_name,
            'usable_variable_technical_file' if add_technical_name else 'usable_variable_file': xlsx_base64
        })

        if download_file:
            url = '/web/content/%s/%s/%s/%s' % (
                self._name, self.id,
                'usable_variable_technical_file' if add_technical_name else 'usable_variable_file',
                self.usable_variable_technical_filename if add_technical_name else self.usable_variable_filename)
            return {'type': 'ir.actions.act_url', 'url': url}

    def _generate_dynamic_fields(self, row, model, sheet, field_type, field, add_technical_name=False):
        return row

    def render_document_report(self, doc_ids, model=False, data=False, output_type=False):
        record = self.env[model or self.model_name].sudo().search([('id', 'in', doc_ids)], limit=1)
        return self.report_id.render_docx_template_data(self, record, data=data, output_type=output_type)

    def _add_print_log(self, docx_file):
        """Adding user printing logs on docx template
        Args:
            docx_file (DocxTemplate): Docx Rendered template
        Returns:
            DocxTemplate: Added footer docx template
        """
        # Printing only if company level setting is allowed
        if not self.env.company.allow_report_print_datetime_log:
            return docx_file

        docx = docx_file.get_docx()
        # Choosing the top most section of the page
        section = docx.sections[0]

        # Calling the footer
        footer = section.footer

        # Calling the paragraph already present in
        # the footer section
        # Set footer paragraph Only if paragraph found in footer
        if footer and footer.paragraphs:
            footer_para = footer.paragraphs[0]

            # Adding text in the footer
            footer_para.text = self.env.company._get_report_footer_print_log()

        if 'normal' in docx.styles:
            footer_style = docx.styles['normal']
            footer_style.paragraph_format.space_after = Pt(12)
            footer_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_style.font.size = Pt(10)

        # Overriding origin docx
        docx_file.docx = docx
        return docx_file
